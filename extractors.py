import os
from pathlib import Path

from dotenv import load_dotenv
from markdown import markdown
from bs4 import BeautifulSoup


load_dotenv()

import fitz
import pandas as pd
import pytesseract
import whisper

from PIL import Image


# ============================================================
# CONFIGURATION
# ============================================================

WHISPER_MODEL = os.getenv("WHISPER_MODEL", "small")

TESSERACT_CMD = os.getenv(
    "TESSERACT_CMD",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

FIRECRAWL_API_KEY = os.getenv("FIRECRAWL_API_KEY")


# ============================================================
# TESSERACT CONFIGURATION
# ============================================================

if Path(TESSERACT_CMD).exists():
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


# ============================================================
# FIRECRAWL
# ============================================================

def get_firecrawl_client():
    """
    Create and return a Firecrawl client.
    """

    if not FIRECRAWL_API_KEY:
        raise ValueError(
            "FIRECRAWL_API_KEY is missing from the .env file."
        )

    from firecrawl import Firecrawl

    return Firecrawl(api_key=FIRECRAWL_API_KEY)


def extract_url(url):
    """
    Extract readable content from a web page using Firecrawl.

    Returns:
        tuple:
            text: extracted markdown text
            ocr_used: always False for URLs
    """

    app = get_firecrawl_client()

    result = app.scrape(
        url,
        formats=["markdown"]
    )

    markdown_text = result.markdown
    html = markdown(markdown_text)

    plain_text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
    print(plain_text)
    print(type(plain_text))

    # Firecrawl SDK response
    if hasattr(result, "markdown"):
        text = result.markdown or ""

    # Fallback if response is a dictionary
    elif isinstance(result, dict):
        text = result.get("markdown", "") or ""

    else:
        text = ""

    return text.strip(), False


# ============================================================
# IMAGE
# ============================================================

def extract_image(path):
    """
    Extract text from an image using Tesseract OCR.
    """

    image = Image.open(path)

    text = pytesseract.image_to_string(image)

    return text.strip(), True


# ============================================================
# PDF
# ============================================================

def extract_pdf(path):
    """
    Extract text from a PDF.

    Normal text PDFs:
        PyMuPDF extracts the text.

    Scanned PDFs:
        Tesseract OCR is used as a fallback.
    """

    parts = []
    ocr_used = False

    doc = fitz.open(path)

    try:
        for i, page in enumerate(doc):

            # ---------------------------------------------
            # Try normal PDF text extraction
            # ---------------------------------------------

            text = page.get_text("text").strip()

            if text:

                parts.append(
                    f"Page {i + 1}\n{text}"
                )

            # ---------------------------------------------
            # OCR fallback
            # ---------------------------------------------

            else:

                pix = page.get_pixmap(
                    matrix=fitz.Matrix(1.5, 1.5),
                    alpha=False
                )

                image = Image.frombytes(
                    "RGB",
                    [pix.width, pix.height],
                    pix.samples
                )

                text = pytesseract.image_to_string(
                    image
                ).strip()

                if text:

                    parts.append(
                        f"Page {i + 1}\n{text}"
                    )

                    ocr_used = True

    finally:
        doc.close()

    return "\n\n".join(parts), ocr_used


# ============================================================
# DOCX
# ============================================================

def extract_docx(path):
    """
    Extract paragraphs and tables from DOCX.
    """

    from docx import Document

    doc = Document(path)

    output = []

    # Paragraphs
    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            output.append(text)

    # Tables
    for table in doc.tables:

        for row in table.rows:

            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
            )

            if row_text:
                output.append(row_text)

    return "\n".join(output)


# ============================================================
# PPTX
# ============================================================

def extract_pptx(path):
    """
    Extract text from PowerPoint slides.
    """

    from pptx import Presentation

    prs = Presentation(path)

    output = []

    for slide_number, slide in enumerate(
        prs.slides,
        start=1
    ):

        texts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    texts.append(text)

        if texts:

            output.append(
                f"Slide {slide_number}\n"
                + "\n".join(texts)
            )

    return "\n\n".join(output)


# ============================================================
# XLSX
# ============================================================

def extract_xlsx(path):
    """
    Extract data from all Excel sheets.
    """

    book = pd.ExcelFile(path)

    output = []

    for sheet in book.sheet_names:

        df = pd.read_excel(
            path,
            sheet_name=sheet,
            dtype=str
        ).fillna("")

        csv_text = df.to_csv(
            index=False
        )

        output.append(
            f"Sheet: {sheet}\n{csv_text}"
        )

    return "\n\n".join(output)


# ============================================================
# CSV
# ============================================================

def extract_csv(path):
    """
    Extract CSV data as text.
    """

    df = pd.read_csv(
        path,
        dtype=str
    ).fillna("")

    return df.to_csv(
        index=False
    )


# ============================================================
# TXT / MARKDOWN
# ============================================================

def extract_text_file(path):
    """
    Extract plain text or Markdown files.
    """

    return path.read_text(
        encoding="utf-8",
        errors="ignore"
    )


# ============================================================
# AUDIO
# ============================================================

def extract_audio(path):
    """
    Transcribe audio using Whisper.
    """

    model = whisper.load_model(
        WHISPER_MODEL
    )

    result = model.transcribe(
        str(path),
        fp16=False
    )

    return result.get(
        "text",
        ""
    ).strip()


# ============================================================
# LOCAL FILE ROUTER
# ============================================================

def extract_file(path):
    """
    Automatically select the correct extractor
    based on the file extension.
    """

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {path}"
        )

    ext = path.suffix.lower()

    # --------------------------------------------------------
    # IMAGE
    # --------------------------------------------------------

    if ext in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".bmp",
        ".tiff"
    }:

        text, ocr = extract_image(path)

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    elif ext == ".pdf":

        text, ocr = extract_pdf(path)
        

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif ext == ".docx":

        text = extract_docx(path)
        ocr = False

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    elif ext == ".pptx":

        text = extract_pptx(path)
        ocr = False

    # --------------------------------------------------------
    # XLSX
    # --------------------------------------------------------

    elif ext == ".xlsx":

        text = extract_xlsx(path)
        ocr = False

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    elif ext == ".csv":

        text = extract_csv(path)
        ocr = False

    # --------------------------------------------------------
    # TXT / MARKDOWN
    # --------------------------------------------------------

    elif ext in {
        ".txt",
        ".md"
    }:

        text = extract_text_file(path)
        ocr = False

    # --------------------------------------------------------
    # AUDIO
    # --------------------------------------------------------

    elif ext in {
        ".wav",
        ".mp3",
        ".m4a",
        ".webm",
        ".ogg",
        ".flac"
    }:

        text = extract_audio(path)
        ocr = False

    else:

        raise ValueError(
            f"Unsupported extension: {ext}"
        )

    return {
        "text": text,
        "ocr": ocr,
        "source_type": "file",
        "source": str(path)
    }


# ============================================================
# URL ROUTER
# ============================================================

def extract_url_content(url):
    """
    Extract content from a URL using Firecrawl.
    """

    text, ocr = extract_url(url)

    return {
        "text": text,
        "ocr": ocr,
        "source_type": "url",
        "source": url
    }


# ============================================================
# UNIVERSAL EXTRACTOR
# ============================================================

def extract_source(source):
    """
    Automatically determine whether the source is
    a local file or a URL.

    Examples:

        extract_source(
            "data/uploads/company.pdf"
        )

        extract_source(
            "https://example.com/company"
        )
    """

    source = str(source).strip()

    # --------------------------------------------------------
    # URL
    # --------------------------------------------------------

    if (
        source.startswith("http://")
        or source.startswith("https://")
    ):

        return extract_url_content(source)

    # --------------------------------------------------------
    # LOCAL FILE
    # --------------------------------------------------------

    return extract_file(
        Path(source)
    )